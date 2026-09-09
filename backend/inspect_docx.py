from docx import Document
from lxml import etree

ns = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
doc = Document('templates/invoice/aca-technologies.docx')

# Count actual tc elements in each row of Table 3
t3 = doc.tables[3]
for ri, row in enumerate(t3.rows):
    tcs = row._tr.findall(f'{{{ns}}}tc')
    texts = []
    for tc in tcs:
        txt = ''.join(t.text or '' for t in tc.iter(f'{{{ns}}}t'))
        texts.append(repr(txt[:40]))
    print(f'T3 Row {ri}: {len(tcs)} cells | {texts}')

print()
t2 = doc.tables[2]
for ri, row in enumerate(t2.rows):
    tcs = row._tr.findall(f'{{{ns}}}tc')
    texts = [repr(''.join(t.text or '' for t in tc.iter(f'{{{ns}}}t'))[:40]) for tc in tcs]
    print(f'T2 Row {ri}: {len(tcs)} cells | {texts}')
