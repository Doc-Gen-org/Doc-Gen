"""
Injects Jinja2 variables into the DOCX invoice template.
Run once after editing the template layout in Word.
"""
from docx import Document
from lxml import etree

NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
XML_SPACE = '{http://www.w3.org/XML/1998/namespace}space'


def _make_run(text):
    r = etree.Element(f'{{{NS}}}r')
    t = etree.SubElement(r, f'{{{NS}}}t')
    t.set(XML_SPACE, 'preserve')
    t.text = text
    return r


def append_value_para(tc, variable_text, align=None):
    """Append a new paragraph containing variable_text to an existing cell."""
    p = etree.SubElement(tc, f'{{{NS}}}p')
    if align:
        pPr = etree.SubElement(p, f'{{{NS}}}pPr')
        jc = etree.SubElement(pPr, f'{{{NS}}}jc')
        jc.set(f'{{{NS}}}val', align)
    p.append(_make_run(variable_text))


def set_cell_value(tc, variable_text):
    """Write variable_text into the first (empty) paragraph of a cell."""
    para = tc.find(f'{{{NS}}}p')
    if para is None:
        para = etree.SubElement(tc, f'{{{NS}}}p')
    for r in para.findall(f'{{{NS}}}r'):
        para.remove(r)
    para.append(_make_run(variable_text))


def inject(docx_path):
    doc = Document(docx_path)

    def tcs(table_idx, row_idx):
        return doc.tables[table_idx].rows[row_idx]._tr.findall(f'{{{NS}}}tc')

    # Table 1 — Trainer Name (row 0) and PO Number (row 1), each label spans full width
    append_value_para(tcs(1, 0)[0], '{{ trainer_name }}')
    append_value_para(tcs(1, 1)[0], '{{ po_number }}')

    # Table 2 — Invoice Number (cell 0) and Invoice Date (cell 1)
    append_value_para(tcs(2, 0)[0], '{{ invoice_number }}')
    append_value_para(tcs(2, 0)[1], '{{ invoice_date }}')

    # Table 4 — Total Amount: label in cell 0, empty value cell is cell 1
    set_cell_value(tcs(4, 0)[1], '{% if total_amount %}{{ total_amount }}{% endif %}')

    # Table 5 — Technology (row 0) and Attend Date (row 1)
    append_value_para(tcs(5, 0)[0], '{{ technology }}')
    append_value_para(tcs(5, 1)[0], '{{ attended_dates_display }}')

    doc.save(docx_path)
    print(f'Saved: {docx_path}')

    # Verify
    doc2 = Document(docx_path)
    print('\nVerification — Jinja2 tags now in template:')
    for ti, table in enumerate(doc2.tables):
        for ri, row in enumerate(table.rows):
            for tc in row._tr.findall(f'{{{NS}}}tc'):
                text = ''.join(t.text or '' for t in tc.iter(f'{{{NS}}}t'))
                if '{{' in text or '{%' in text:
                    label = text.split('\n')[0][:40]
                    tag = [l.strip() for l in text.split('\n') if '{{' in l or '{%' in l]
                    print(f'  Table {ti} Row {ri}: {repr(label)} → {tag}')


if __name__ == '__main__':
    inject('templates/invoice/aca-technologies.docx')
