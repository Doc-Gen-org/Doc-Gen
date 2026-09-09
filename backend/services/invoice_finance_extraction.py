import os
import re
import shutil
from docx import Document
from docx.oxml.ns import qn

# Word content-control tag -> canonical finance label
_TAG_TO_LABEL = {
    "account_holder_name": "Account Holder Name",
    "bank_name": "Bank Name",
    "account_number": "Account Number",
    "pan_number": "PAN Number",
    "ifsc_code": "IFSC Code",
    "mode_of_payment": "Mode of Payment",
}

# Plain-cell label text (lowercased) -> canonical finance label
_LABEL_TO_LABEL = {
    "account holder name": "Account Holder Name",
    "bank name": "Bank Name",
    "account number": "Account Number",
    "account name": "Account Number",
    "pan number": "PAN Number",
    "ifsc code": "IFSC Code",
    "mode of payment": "Mode of Payment",
}

_PLACEHOLDER_TEXTS = {"click or tap here to enter text."}


def _clean_value(text: str) -> str:
    """Normalise a cell/control value, discarding placeholders and unrendered tokens."""
    value = (text or "").strip()
    if not value:
        return ""
    if value.lower() in _PLACEHOLDER_TEXTS:
        return ""
    if value.startswith("{{") or value.startswith("{%"):
        return ""
    return value


def _element_text(element) -> str:
    return "".join(t.text or "" for t in element.iter(qn("w:t")))


def _row_slots(tr) -> list[dict]:
    """
    Ordered list of value-bearing slots in a table row. Each slot is a plain
    cell (<w:tc>) or a content control (<w:sdt>, which python-docx's row.cells
    skips). Returns {"text": visible_text, "tag": sdt_tag_or_None}.
    """
    slots = []
    for child in tr:
        name = child.tag.split("}")[-1]
        if name == "tc":
            slots.append({"text": _element_text(child), "tag": None, "el": child})
        elif name == "sdt":
            pr = child.find(qn("w:sdtPr"))
            tag_el = pr.find(qn("w:tag")) if pr is not None else None
            tag = tag_el.get(qn("w:val")) if tag_el is not None else None
            slots.append({"text": _element_text(child), "tag": tag, "el": child})
    return slots


def extract_finance_details_from_docx(file_path: str) -> dict:
    """
    Extracts finance fields from a filled invoice DOCX. Handles the current
    template format where values live either in plain cells next to their label
    or inside Word content controls (SDTs) tagged with the field name.
    """
    doc = Document(file_path)
    results = {}

    for table in doc.tables:
        for tr in table._tbl.findall(qn("w:tr")):
            slots = _row_slots(tr)
            for i, slot in enumerate(slots):
                # Content control: its tag names the field directly
                if slot["tag"] in _TAG_TO_LABEL:
                    value = _clean_value(slot["text"])
                    if value:
                        results[_TAG_TO_LABEL[slot["tag"]]] = value
                    continue

                # Plain label cell: the value is the following slot
                canonical = _LABEL_TO_LABEL.get(slot["text"].strip().lower())
                if canonical and i + 1 < len(slots):
                    value = _clean_value(slots[i + 1]["text"])
                    if value:
                        results[canonical] = value

    return results


# Ordered list of (canonical_label, regex) — most specific patterns first
_PDF_FIELD_PATTERNS = [
    ("Account Holder Name", r"Account\s+Holder\s+Name\s+(.+?)(?=\n|Bank\s+Name|Account\s+Number|Account\s+Name|PAN|IFSC|Mode)"),
    ("Bank Name",           r"Bank\s+Name\s+(.+?)(?=\n|Account\s+Number|Account\s+Name|PAN|IFSC|Mode)"),
    ("Account Number",      r"Account\s+(?:Number|Name)\s+([A-Za-z0-9]+)"),
    ("PAN Number",          r"PAN\s+Number\s+([A-Z]{5}[0-9]{4}[A-Z])"),
    ("IFSC Code",           r"IFSC\s+Code\s+([A-Z]{4}0[A-Z0-9]{6})"),
    ("Mode of Payment",     r"Mode\s+of\s+Payment\s+(.+?)(?:\n|$)"),
    ("Training Days",       r"Training\s+Days\s*(?:\(Qty\))?\s+(\d+)"),
    ("Rate / Day",          r"Rate\s*[/\\]\s*Day\s+([\d,]+(?:\.\d{1,2})?)"),
]


def extract_finance_details_from_pdf(file_path: str) -> dict:
    """
    Extracts known finance fields from a filled PDF invoice using targeted
    regex patterns against the full page text. Works for digitally-filled
    or Word-to-PDF converted invoices; scanned images won't extract.
    """
    try:
        import pdfplumber
        import re
    except ImportError:
        return {}

    results = {}
    try:
        with pdfplumber.open(file_path) as pdf:
            full_text = "\n".join(page.extract_text() or "" for page in pdf.pages)

        for label, pattern in _PDF_FIELD_PATTERNS:
            m = re.search(pattern, full_text, re.IGNORECASE | re.DOTALL)
            if m:
                value = m.group(1).strip()
                if value:
                    results[label] = value

    except Exception:
        return {}

    return results


def extract_finance_details(file_path: str) -> dict:
    extension = os.path.splitext(file_path)[1].lower()
    if extension == ".docx":
        return extract_finance_details_from_docx(file_path)
    elif extension == ".pdf":
        return extract_finance_details_from_pdf(file_path)
    return {}

# Map of label text (lowercase) as it appears in the DOCX → account field name
_BANK_LABEL_MAP = {
    "account holder name": "account_holder_name",
    "bank name": "bank_name",
    "account name": "account_number",   # DOCX template uses "Account Name"
    "account number": "account_number",
    "pan number": "pan_number",
    "ifsc code": "ifsc_code",
    "mode of payment": "mode_of_payment",
}


def _set_cell_text(cell, text: str) -> None:
    """Replace all text in a table cell while preserving the first run's formatting."""
    for para in cell.paragraphs:
        for run in para.runs:
            run.text = ""
    if not cell.paragraphs:
        return
    para = cell.paragraphs[0]
    if para.runs:
        para.runs[0].text = text
    else:
        para.add_run(text)


def replace_bank_details_in_docx(source_path: str, dest_path: str, account) -> str:
    """
    Copy source_path → dest_path, then replace every bank/payment detail
    cell in the DOCX tables with the values from `account` (a CompanyAccount).
    Handles 2-cell rows (label | value) and 4-cell rows (two pairs side-by-side).
    """
    shutil.copy2(source_path, dest_path)
    doc = Document(dest_path)

    for table in doc.tables:
        for row in table.rows:
            cells = row.cells
            # Build index of (label_cell_index, value_cell_index) pairs for this row
            pairs = []
            if len(cells) == 2:
                pairs = [(0, 1)]
            elif len(cells) == 4:
                pairs = [(0, 1), (2, 3)]
            else:
                # For rows with more cells, scan each cell as a potential label
                for i in range(len(cells) - 1):
                    pairs.append((i, i + 1))

            for label_idx, value_idx in pairs:
                label = cells[label_idx].text.strip().lower()
                field = _BANK_LABEL_MAP.get(label)
                if field:
                    new_value = getattr(account, field, "") or ""
                    _set_cell_text(cells[value_idx], new_value)

    doc.save(dest_path)
    return dest_path


# Content-control tags used for bank/payment fields in the invoice DOCX
_BANK_FIELD_TAGS = {
    "account_holder_name",
    "bank_name",
    "account_number",
    "pan_number",
    "ifsc_code",
    "mode_of_payment",
}


def _parse_amount(text) -> float | None:
    """Extract the first numeric value from arbitrary text (strips commas/currency)."""
    if text is None:
        return None
    m = re.search(r"-?\d+(?:\.\d+)?", str(text).replace(",", ""))
    if not m:
        return None
    try:
        return float(m.group(0))
    except ValueError:
        return None


def _fmt_amount(value) -> str:
    """Render a number as an invoice amount: integer when whole, else 2dp, with thousands separators."""
    if value is None:
        return ""
    try:
        f = float(value)
    except (TypeError, ValueError):
        return str(value)
    if f == int(f):
        return f"{int(f):,}"
    return f"{f:,.2f}"


def _set_element_text(element, text: str) -> None:
    """Set the visible text of a tc/sdt element, preserving the first run's formatting."""
    t_list = list(element.iter(qn("w:t")))
    if not t_list:
        return
    t_list[0].text = text
    for extra in t_list[1:]:
        extra.text = ""


def extract_invoice_amounts_from_pdf(file_path: str) -> dict:
    """
    Best-effort parse of the invoice items table from a PDF: the training
    quantity (days) and each reimbursement amount (travel / food /
    accommodation). Amounts are taken as the last number on each item line.
    """
    try:
        import pdfplumber
    except ImportError:
        return {}

    result: dict = {}
    try:
        with pdfplumber.open(file_path) as pdf:
            full_text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    except Exception:
        return {}

    attend = re.search(r"Attend\s+Date\s+(.+)", full_text, re.IGNORECASE)
    if attend:
        value = attend.group(1).strip()
        if value:
            result["attend_date"] = value

    for line in full_text.splitlines():
        m = re.match(r"\s*\d+\s+(.*)", line)
        if not m:
            continue
        desc = m.group(1)
        nums = re.findall(r"\d[\d,]*(?:\.\d+)?", desc)
        if not nums:
            continue
        amount = _parse_amount(nums[-1])
        dlow = desc.lower()
        if "total" in dlow:
            continue
        if "training" in dlow or "remuneration" in dlow:
            if len(nums) >= 3:
                result["training_days"] = _parse_amount(nums[0])
        elif "travel" in dlow:
            result["travel_amount"] = amount
        elif "food" in dlow:
            result["food_amount"] = amount
        elif "accommodation" in dlow or "stay" in dlow or "lodg" in dlow:
            result["accommodation_amount"] = amount

    return result


def modify_uploaded_invoice_docx(source_path: str, dest_path: str, account, rate_per_day) -> str:
    """
    Copy source_path → dest_path, then edit the filled invoice DOCX in place:
      1. Replace bank/payment detail cells with the values from `account`.
      2. Replace the training row's Rate / Day with `rate_per_day`.
      3. Recompute the training amount (days × rate) and the total
         (training amount + all reimbursement amounts).
    Handles both plain table cells and Word content controls (SDTs).
    """
    shutil.copy2(source_path, dest_path)
    doc = Document(dest_path)
    rate = _parse_amount(rate_per_day) or 0.0

    # --- 1. Bank / payment details (SDT-aware) ---
    for table in doc.tables:
        for tr in table._tbl.findall(qn("w:tr")):
            slots = _row_slots(tr)
            for i, slot in enumerate(slots):
                if slot.get("tag") in _BANK_FIELD_TAGS:
                    _set_element_text(slot["el"], getattr(account, slot["tag"], "") or "")
                    continue
                field = _BANK_LABEL_MAP.get(slot["text"].strip().lower())
                if field and i + 1 < len(slots):
                    _set_element_text(slots[i + 1]["el"], getattr(account, field, "") or "")

    # --- 2 & 3. Items table: rate, training amount, and total ---
    running_total = 0.0
    for table in doc.tables:
        rows = table._tbl.findall(qn("w:tr"))
        if not rows:
            continue
        header_texts = [s["text"].strip().lower() for s in _row_slots(rows[0])]
        if not any("rate" in t for t in header_texts) or not any("amount" in t for t in header_texts):
            continue  # not the items table

        def _col(*keywords):
            for idx, t in enumerate(header_texts):
                if any(k in t for k in keywords):
                    return idx
            return None

        rate_col = _col("rate")
        amount_col = _col("amount")
        qty_col = _col("days", "qty")
        desc_col = _col("description", "service")

        for tr in rows[1:]:
            slots = _row_slots(tr)
            if amount_col is None or amount_col >= len(slots):
                continue
            joined = " ".join(s["text"].lower() for s in slots)
            if "total" in joined:
                continue  # total row handled separately
            desc_text = slots[desc_col]["text"].lower() if (desc_col is not None and desc_col < len(slots)) else ""
            is_training = "training" in desc_text or "remuneration" in desc_text
            if is_training and rate_col is not None and rate_col < len(slots):
                qty = _parse_amount(slots[qty_col]["text"]) if (qty_col is not None and qty_col < len(slots)) else None
                _set_element_text(slots[rate_col]["el"], _fmt_amount(rate))
                if qty is not None:
                    amt = qty * rate
                    _set_element_text(slots[amount_col]["el"], _fmt_amount(amt))
                    running_total += amt
                    continue
            existing = _parse_amount(slots[amount_col]["text"])
            if existing is not None:
                running_total += existing

    # Total row may live in the items table footer or a separate table
    for table in doc.tables:
        for tr in table._tbl.findall(qn("w:tr")):
            slots = _row_slots(tr)
            if not slots:
                continue
            if "total" in " ".join(s["text"].lower() for s in slots):
                _set_element_text(slots[-1]["el"], _fmt_amount(running_total))

    doc.save(dest_path)
    return dest_path